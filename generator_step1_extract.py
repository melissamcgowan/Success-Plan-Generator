"""
Data extraction layer for the Success Plan Generator.

Today this reads the synthetic workbook (data/success_plan_source_data.xlsx),
built to mirror a typical CSP export (Gainsight / Totango / ChurnZero).
When a real CSP connection is available, only this file changes — swap the
load_workbook() calls below for API calls that return the same field names,
and build_deck.js (the rendering layer) needs no changes at all.
"""
import json
import sys
import openpyxl
from datetime import datetime, date

SRC = "/home/claude/success_plan/data/success_plan_source_data.xlsx"


def rows_for(ws, header, account_id, id_col="account_id"):
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx[id_col]] == account_id:
            out.append({name: row[idx[name]] for name in header})
    return out


def fmt_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%b %-d, %Y")
    if isinstance(v, str):
        try:
            return datetime.strptime(v, "%Y-%m-%d").strftime("%b %-d, %Y")
        except ValueError:
            return v
    return str(v)


def build_payload(account_id):
    wb = openpyxl.load_workbook(SRC, data_only=True)

    accounts_ws = wb["Accounts"]
    acc_header = [c.value for c in next(accounts_ws.iter_rows(min_row=1, max_row=1))]
    account = rows_for(accounts_ws, acc_header, account_id)[0]

    def sheet_rows(name):
        ws = wb[name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        return rows_for(ws, header, account_id)

    stakeholders = sheet_rows("Stakeholders")
    goals_uc = sheet_rows("Goals & Use Cases")
    usage = sheet_rows("Usage & Adoption")
    milestones = sheet_rows("Milestones")
    metrics = sheet_rows("Success Metrics")
    actions = sheet_rows("Actions")
    risks = sheet_rows("Risks")

    payload = {
        "generated": datetime.now().strftime("%B %-d, %Y"),
        "account": {
            **account,
            "renewal_date_fmt": fmt_date(account["renewal_date"]),
            "arr_fmt": f"${account['arr_usd']:,.0f}",
        },
        "stakeholders": stakeholders,
        "goals": [g for g in goals_uc if g["type"] == "Business Goal"],
        "use_cases": [g for g in goals_uc if g["type"] == "Use Case"],
        "usage": usage,
        "milestones": [
            {**m, "target_date_fmt": fmt_date(m["target_date"])} for m in milestones
        ],
        "metrics": metrics,
        "actions": [
            {**a, "due_date_fmt": fmt_date(a["due_date"])} for a in actions
        ],
        "risks": risks,
    }
    return payload


if __name__ == "__main__":
    account_id = sys.argv[1] if len(sys.argv) > 1 else "ACC-1001"
    payload = build_payload(account_id)
    out_path = f"/home/claude/success_plan/data/payload_{account_id}.json"
    with open(out_path, "w") as f:
        json.dump(payload, f, indent=2, default=str)
    print("saved", out_path)
